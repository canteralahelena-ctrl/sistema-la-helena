import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def read_csv(path):
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def to_number(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return 0.0


def write_table(sheet, row, col, headers, rows, fill):
    white = "FFFFFF"
    border = Side(style="thin", color="CBD5E1")
    hair = Side(style="hair", color="E2E8F0")
    for offset, header in enumerate(headers):
        cell = sheet.cell(row, col + offset, header)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=border)

    for r_idx, values in enumerate(rows, row + 1):
        for c_idx, value in enumerate(values, col):
            cell = sheet.cell(r_idx, c_idx, value)
            cell.border = Border(bottom=hair)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--resumen", required=True)
    parser.add_argument("--fuentes", required=True)
    parser.add_argument("--detalle", required=True)
    parser.add_argument("--no-clasificados", required=True)
    parser.add_argument("--salida", required=True)
    args = parser.parse_args()

    resumen = read_csv(args.resumen)
    fuentes = read_csv(args.fuentes)
    detalle = read_csv(args.detalle)
    no_clasificados = read_csv(args.no_clasificados)

    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    sh_fuentes = wb.create_sheet("FUENTES")
    sh_detalle = wb.create_sheet("DETALLE")
    sh_no = wb.create_sheet("NO_CLASIFICADOS")

    dark = "334155"
    blue = "1D4ED8"
    orange = "C2410C"
    gray = "F1F5F9"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False

    ws["A1"] = "Analisis comercial por categorias"
    ws["A1"].font = Font(bold=True, size=16, color="1F2937")
    ws["A1"].fill = PatternFill("solid", fgColor=gray)
    ws.merge_cells("A1:J1")
    ws["A2"] = (
        "El reporte principal incluye solo categorias comerciales validas. "
        "Los productos o servicios no clasificados se excluyen y se listan en NO_CLASIFICADOS."
    )
    ws["A2"].font = Font(color="475569")
    ws.merge_cells("A2:J2")

    resumen_headers = list(resumen[0].keys()) if resumen else []
    resumen_rows = []
    for row in resumen:
        resumen_rows.append(
            [
                row.get("Periodo", ""),
                row.get("LineaNegocio", ""),
                to_number(row.get("Importe")),
                to_number(row.get("ParticipacionPct")) / 100,
                to_number(row.get("CantidadTotal")),
                to_number(row.get("M3Estimado")),
                int(to_number(row.get("CantidadComprobantes"))),
                int(to_number(row.get("CantidadClientes"))),
                int(to_number(row.get("CantidadOperaciones"))),
                row.get("Fuentes", ""),
            ]
        )
    resumen_headers = [
        "Periodo",
        "LineaNegocio",
        "Importe",
        "Participacion",
        "CantidadTotal",
        "M3Estimado",
        "CantidadComprobantes",
        "CantidadClientes",
        "CantidadOperaciones",
        "Fuentes",
    ]
    write_table(ws, 4, 1, resumen_headers, resumen_rows, blue)
    for excel_row in range(5, 5 + len(resumen_rows)):
        ws.cell(excel_row, 4).number_format = "0.00%"

    if resumen_rows:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Categorias comerciales validas"
        chart.y_axis.title = "Importe"
        chart.x_axis.title = "Categoria"
        chart.height = 11
        chart.width = 24
        chart.legend = None
        values = Reference(ws, min_col=3, min_row=4, max_row=4 + len(resumen_rows))
        cats = Reference(ws, min_col=2, min_row=5, max_row=4 + len(resumen_rows))
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.numFmt = "$ #,##0"
        ws.add_chart(chart, "L4")

    total_no = sum(to_number(row.get("ImporteSinIVA")) for row in no_clasificados)
    ws["A16"] = (
        f"Existen {len(no_clasificados)} productos/servicios no clasificados "
        f"por un total de ${total_no:,.2f}. Revisar hoja NO_CLASIFICADOS."
    )
    ws["A16"].font = Font(bold=True, color=orange)
    ws.merge_cells("A16:J16")

    fuente_headers = list(fuentes[0].keys()) if fuentes else []
    fuente_rows = [[to_number(v) if k in {"Importe", "CantidadComprobantes", "CantidadClientes", "CantidadOperaciones"} else v for k, v in row.items()] for row in fuentes]
    if fuente_headers:
        write_table(sh_fuentes, 1, 1, fuente_headers, fuente_rows, dark)

    detalle_headers = list(detalle[0].keys()) if detalle else []
    detalle_rows = []
    for row in detalle:
        detalle_rows.append([to_number(v) if k in {"Cantidad", "M3Estimado", "Importe"} else v for k, v in row.items()])
    if detalle_headers:
        write_table(sh_detalle, 1, 1, detalle_headers, detalle_rows, dark)

    no_headers = [
        "Fecha",
        "TipoComprobante",
        "NumeroComprobante",
        "NumeroFactura",
        "NumeroRMT",
        "Cliente",
        "ProductoDescripcionOriginal",
        "Cantidad",
        "Unidad",
        "ImporteSinIVA",
        "CategoriaSugerida",
        "NivelConfianza",
        "MotivoNoClasificacion",
        "ObservacionManual",
    ]
    no_rows = []
    for row in no_clasificados:
        no_rows.append(
            [
                row.get("Fecha", ""),
                row.get("TipoComprobante", ""),
                row.get("NumeroComprobante", ""),
                row.get("NumeroFactura", ""),
                row.get("NumeroRMT", ""),
                row.get("Cliente", ""),
                row.get("ProductoDescripcionOriginal", ""),
                to_number(row.get("Cantidad")),
                row.get("Unidad", ""),
                to_number(row.get("ImporteSinIVA")),
                row.get("CategoriaSugerida", ""),
                row.get("NivelConfianza", ""),
                row.get("MotivoNoClasificacion", ""),
                row.get("ObservacionManual", ""),
            ]
        )
    write_table(sh_no, 1, 1, no_headers, no_rows, orange)

    widths = {
        "RESUMEN": [14, 28, 16, 16, 16, 14, 20, 18, 20, 18],
        "FUENTES": [14, 28, 14, 16, 20, 18, 20],
        "DETALLE": [14, 14, 12, 14, 14, 14, 20, 44, 24, 16, 12, 12, 14, 14, 16, 12, 60],
        "NO_CLASIFICADOS": [14, 16, 18, 18, 18, 30, 48, 12, 12, 16, 22, 16, 48, 28],
    }
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for idx, width in enumerate(widths.get(sheet.title, []), 1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    out = Path(args.salida)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    check = load_workbook(out, read_only=False, data_only=False)
    assert "NO_CLASIFICADOS" in check.sheetnames


if __name__ == "__main__":
    main()
