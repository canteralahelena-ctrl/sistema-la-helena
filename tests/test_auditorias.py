import csv
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


class AuditoriasTests(unittest.TestCase):
    def test_reglas_anulados_en_powershell(self):
        result = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(ROOT / "test_auditoria_anulados.ps1"),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        self.assertIn("FTS A 200004460", result.stdout)

    def test_reglas_precio_venta_bajo_en_powershell(self):
        result = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(ROOT / "test_auditoria_precio_venta.ps1"),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        self.assertIn("Caso RMT 30715 detectado", result.stdout)

    def test_resolucion_controlada_detalle_en_powershell(self):
        result = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(ROOT / "test_auditoria_detalles.ps1"),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        self.assertIn("TIPO+Numero unico legacy permitido", result.stdout)

    def test_localizador_historico_y_rmt_en_powershell(self):
        result = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(ROOT / "test_auditoria_documentos.ps1"),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        self.assertIn("RMT_030754 encontrado", result.stdout)

    def test_motor_compartido_cubre_los_tres_usuarios(self):
        orchestrator = (ROOT / "auditoria_semanal_maxi.ps1").read_text(encoding="utf-8-sig")
        engine = (ROOT / "auditoria_usuario_3.ps1").read_text(encoding="utf-8-sig")
        for user_id in (1, 2, 3):
            self.assertIn(f"Id = {user_id}", orchestrator)
        self.assertIn('"auditoria_usuario_3.ps1"', orchestrator)
        self.assertIn('"-UsuarioId",', orchestrator)
        self.assertIn("$User.Id", orchestrator)
        self.assertIn('"auditoria_anulados.ps1"', engine)
        self.assertIn("T_GAS_NotasCreditoAplicaciones", engine)
        self.assertIn("Get-AuditComprobantesAnulados", engine)
        self.assertIn("Resolve-AuditComprobanteSaldo", engine)
        self.assertIn('"auditoria_documentos.ps1"', engine)
        self.assertIn("Find-AuditDocumentPdf", engine)
        self.assertIn("NO_COMPARADO_ANULADA", engine)
        self.assertNotIn('Join-Path $root "work\\cache"', engine)
        self.assertIn("HELENA_PYTHON_EXECUTABLE", orchestrator)
        self.assertNotIn("codex-runtimes", orchestrator)

    def test_motor_es_solo_lectura(self):
        engine = (ROOT / "auditoria_usuario_3.ps1").read_text(encoding="utf-8-sig").upper()
        for token in ("INSERT INTO", "UPDATE ", "DELETE FROM", "DROP TABLE", "ALTER TABLE"):
            self.assertNotIn(token, engine)

    def test_excel_agrega_campos_sin_cambiar_presentacion(self):
        with tempfile.TemporaryDirectory() as tmp:
            sandbox = Path(tmp) / "work_v2"
            outputs = Path(tmp) / "outputs"
            sandbox.mkdir()
            outputs.mkdir()
            shutil.copy2(ROOT / "generar_excel_auditoria.py", sandbox / "generar_excel_auditoria.py")

            prefix = "fixture"
            suffix = "2026"
            comprobantes_headers = [
                "Fecha", "Tipo", "Numero", "IdCOMPROVANTE", "Cliente", "Importe", "Saldo",
                "Estado", "Observacion", "NotaCreditoRelacionada", "PdfEstado", "PdfEsperado",
                "PdfTotal", "PdfTotalEstado", "PdfDiferencia", "Alerta",
            ]
            fixtures = {
                "comprobantes": (comprobantes_headers, [
                    "30/07/2026", "FTS A", "200004460", "44603", "CLIENTE", "562650", "0",
                    "ANULADA CON NC", "VERIFICAR POR QUÉ FUE ANULADA",
                    "NCS A 00002-00000125", "ENCONTRADO", "factura.pdf", "",
                    "NO_COMPARADO_ANULADA", "", "ANULADA CON NC",
                ]),
                "notas_credito": (["Fecha", "TipoNC", "NumeroNC"], ["30/07/2026", "NCS A", "125"]),
                "pagos": (["Fecha", "IdPAGO", "Alerta"], ["30/07/2026", "1", ""]),
                "entregas": (["Fecha", "IdPAGO", "Importe"], ["30/07/2026", "1", "10"]),
            }
            for name, (headers, row) in fixtures.items():
                path = outputs / f"{prefix}_{name}_{suffix}.csv"
                with path.open("w", newline="", encoding="utf-8-sig") as stream:
                    writer = csv.writer(stream)
                    writer.writerow(headers)
                    writer.writerow(row)

            result = subprocess.run(
                [sys.executable, str(sandbox / "generar_excel_auditoria.py"), "--prefix", prefix, "--suffix", suffix],
                cwd=sandbox,
                env={**os.environ, "HELENA_OUTPUTS_DIR": str(outputs)},
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            workbook = load_workbook(outputs / f"{prefix}_{suffix}.xlsx")
            self.assertEqual(["Comprobantes", "Notas de credito", "Pagos", "Medios de pago"], workbook.sheetnames)
            sheet = workbook["Comprobantes"]
            headers = [cell.value for cell in sheet[1]]
            for header in ("Estado", "Observacion", "NotaCreditoRelacionada"):
                self.assertEqual(1, headers.count(header))
            self.assertEqual("A2", sheet.freeze_panes)
            self.assertTrue(sheet.auto_filter.ref)
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual("1F2937", sheet["A1"].fill.fgColor.rgb[-6:])
            self.assertTrue(sheet["A1"].font.bold)


if __name__ == "__main__":
    unittest.main()
