import argparse
import csv
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get("HELENA_OUTPUTS_DIR", Path(__file__).resolve().parent / "outputs"))
parser = argparse.ArgumentParser(description="Genera Excel de auditoria de usuario Access.")
parser.add_argument("--prefix", default="auditoria_usuario_3_maxi")
parser.add_argument("--suffix", default="")
parser.add_argument("--width-template", default="")
args = parser.parse_args()

PREFIX = args.prefix
if args.suffix:
    SUFFIX = args.suffix
else:
    latest_comprobantes = sorted(
        OUT.glob(f"{PREFIX}_comprobantes_*.csv"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not latest_comprobantes and PREFIX == "auditoria_usuario_3_maxi":
        PREFIX = "auditoria_usuario_3"
        latest_comprobantes = sorted(
            OUT.glob(f"{PREFIX}_comprobantes_*.csv"),
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
    if not latest_comprobantes:
        raise SystemExit("No hay CSV de auditoria para generar Excel.")
    SUFFIX = latest_comprobantes[0].stem.replace(f"{PREFIX}_comprobantes_", "")

SHEETS = [
    ("Comprobantes", OUT / f"{PREFIX}_comprobantes_{SUFFIX}.csv"),
    ("Notas de credito", OUT / f"{PREFIX}_notas_credito_{SUFFIX}.csv"),
    ("Pagos", OUT / f"{PREFIX}_pagos_{SUFFIX}.csv"),
    ("Medios de pago", OUT / f"{PREFIX}_entregas_{SUFFIX}.csv"),
]
xlsx_path = OUT / f"{PREFIX}_{SUFFIX}.xlsx"
OUT.mkdir(parents=True, exist_ok=True)
DEFAULT_HEADERS = {
    "Comprobantes": [
        "Fecha", "Tipo", "Numero", "IdCOMPROVANTE", "Cliente", "Importe", "Saldo",
        "Estado", "Observacion", "NotaCreditoRelacionada",
        "PdfEstado", "PdfEsperado", "PdfTotal", "PdfTotalEstado", "PdfDiferencia", "Alerta",
    ],
    "Notas de credito": [
        "Fecha", "TipoNC", "NumeroNC", "IdCOMPROVANTE", "Cliente", "CUIT", "Importe", "IVA",
        "ComprobanteReferenciado", "ReferenciaTextoPDF", "FacturaExiste", "MismoClienteOCuit",
        "PuntoVentaNumero", "IvaCoherente", "TipoAjuste", "ReferenciaDuplicada", "Resultado",
        "PdfEstado", "PdfArchivo",
    ],
    "Pagos": [
        "Fecha", "IdPAGO", "IdCLIENTE", "Cliente", "TipoPagoSistema", "MedioPago", "Monto",
        "Saldo", "ImputadoADocumentos", "TotalMediosPago", "CantidadImputaciones",
        "CantidadMediosPago", "EstadoImputacion", "UserID", "Alerta",
    ],
    "Medios de pago": [
        "Fecha", "IdPAGO", "IdENTREGA", "Cliente", "TipoPagoSistema", "MedioPago",
        "MedioPagoOriginal", "Banco", "Numero", "FechaACobrar", "Importe", "Estado",
        "DepositadoEn", "UserID",
    ],
}

red_fill = PatternFill("solid", fgColor="F8D7DA")
red_font = Font(color="9C0006", bold=True)
yellow_fill = PatternFill("solid", fgColor="FFF3CD")
yellow_font = Font(color="7A4B00", bold=True)
green_fill = PatternFill("solid", fgColor="D1E7DD")
header_fill = PatternFill("solid", fgColor="1F2937")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def load_width_template(path):
    if not path:
        return {}
    template_path = Path(path)
    if not template_path.exists():
        raise SystemExit(f"No existe plantilla de anchos: {template_path}")
    wb = load_workbook(template_path)
    widths = {}
    for ws in wb.worksheets:
        sheet_widths = {}
        for idx, cell in enumerate(ws[1], start=1):
            sheet_widths[idx] = ws.column_dimensions[cell.column_letter].width
        widths[ws.title] = sheet_widths
    return widths


def read_rows(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def style_sheet(ws, headers, width_template=None):
    width_template = width_template or {}
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    alert_col = headers.index("Alerta") + 1 if "Alerta" in headers else None
    pdf_estado_col = headers.index("PdfEstado") + 1 if "PdfEstado" in headers else None
    pdf_total_estado_col = headers.index("PdfTotalEstado") + 1 if "PdfTotalEstado" in headers else None
    estado_imputacion_col = headers.index("EstadoImputacion") + 1 if "EstadoImputacion" in headers else None
    resultado_col = headers.index("Resultado") + 1 if "Resultado" in headers else None

    for row_idx in range(2, ws.max_row + 1):
        alerta = ws.cell(row_idx, alert_col).value if alert_col else ""
        pdf_estado = ws.cell(row_idx, pdf_estado_col).value if pdf_estado_col else ""
        pdf_total_estado = ws.cell(row_idx, pdf_total_estado_col).value if pdf_total_estado_col else ""
        estado_imputacion = ws.cell(row_idx, estado_imputacion_col).value if estado_imputacion_col else ""
        resultado = ws.cell(row_idx, resultado_col).value if resultado_col else ""

        fill = None
        font = None
        if alerta or str(estado_imputacion).startswith("REVISAR"):
            fill = red_fill
            font = red_font
        elif "⚠️" in str(resultado):
            fill = yellow_fill
            font = yellow_font
        elif pdf_estado == "FALTANTE" or str(pdf_total_estado).startswith("NO_LEIDO"):
            fill = yellow_fill
            font = yellow_font
        elif pdf_estado == "ENCONTRADO" and pdf_total_estado in ("OK", ""):
            fill = green_fill

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if fill:
                cell.fill = fill
            if font:
                cell.font = font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for col_idx, header in enumerate(headers, start=1):
        if col_idx in width_template and width_template[col_idx]:
            width = width_template[col_idx]
        else:
            values = [str(ws.cell(row_idx, col_idx).value or "") for row_idx in range(1, min(ws.max_row, 200) + 1)]
            width = max(len(header), *(len(v) for v in values)) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width, 10), 55)


wb = Workbook()
wb.remove(wb.active)
width_templates = load_width_template(args.width_template)

for title, path in SHEETS:
    ws = wb.create_sheet(title=title)
    rows = read_rows(path)
    headers = list(rows[0].keys()) if rows else DEFAULT_HEADERS.get(title, [])
    if not headers:
        ws.append(["Sin datos"])
        continue
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws, headers, width_templates.get(title))

wb.save(xlsx_path)
print(xlsx_path)
