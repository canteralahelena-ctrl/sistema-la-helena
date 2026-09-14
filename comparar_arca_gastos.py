import argparse
import csv
import json
import re
import subprocess
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
PYTHON = Path(r"C:\USUARIO_EJEMPLO\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe")


def fold(text):
    return "".join(
        c for c in unicodedata.normalize("NFD", str(text or "")) if unicodedata.category(c) != "Mn"
    ).lower()


def money(value):
    return f"$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def n(value):
    if value in (None, ""):
        return 0.0
    return float(value)


def is_nc(tipo):
    t = fold(tipo)
    return "nota de credito" in t or t.strip().startswith("nc")


def arca_type(tipo):
    t = fold(tipo)
    if "factura a" in t:
        return "FTA"
    if "factura b" in t:
        return "FTB"
    if "factura c" in t:
        return "FTC"
    if "nota de credito" in t:
        return "NC"
    return str(tipo or "").strip().upper()


def access_type(tipo):
    t = fold(tipo).replace(" ", "")
    if t in {"fta", "fta"}:
        return "FTA"
    if t in {"ftb", "ftb"}:
        return "FTB"
    if t in {"ftc", "ftc"}:
        return "FTC"
    if t.startswith("nc"):
        return "NC"
    return str(tipo or "").strip().upper()


def parse_date(value):
    if isinstance(value, datetime):
        return value
    return datetime.strptime(str(value), "%d/%m/%Y")


def load_arca(path, year, month):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[2]]
    rows = []
    for values in ws.iter_rows(min_row=3, values_only=True):
        if not any(v is not None for v in values):
            continue
        d = dict(zip(headers, values))
        fecha = parse_date(d["Fecha"])
        if fecha.year != year or fecha.month != month:
            continue
        sign = -1 if is_nc(d["Tipo"]) else 1
        tipo_norm = arca_type(d["Tipo"])
        numero = int(d["Número Desde"])
        tipo_cambio = n(d.get("Tipo Cambio") or 1) or 1.0
        moneda = str(d.get("Moneda") or "").strip()
        total_iva = n(d["Total IVA"]) * tipo_cambio
        importe_total = n(d["Imp. Total"]) * tipo_cambio
        row = {
            "origen": "ARCA",
            "fecha": fecha.strftime("%Y-%m-%d"),
            "tipo": d["Tipo"],
            "tipo_norm": tipo_norm,
            "punto_venta": int(d["Punto de Venta"]),
            "numero": numero,
            "numero_texto": str(numero),
            "emisor": d["Denominación Emisor"],
            "cuit": str(d["Nro. Doc. Emisor"]),
            "moneda": moneda,
            "tipo_cambio": tipo_cambio,
            "iva": sign * total_iva,
            "iva_original": n(d["Total IVA"]),
            "importe": sign * importe_total,
            "importe_original": n(d["Imp. Total"]),
        }
        rows.append(row)
    return rows


def run_access_query(year, month):
    script = ROOT / "work" / "exportar_gastos_mes.ps1"
    completed = subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-Anio",
            str(year),
            "-Mes",
            str(month),
        ],
        cwd=str(ROOT),
        text=True,
        capture_output=True,
        timeout=120,
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr or completed.stdout)
    return json.loads(completed.stdout)


def access_to_rows(data):
    rows = []
    for d in data:
        tipo_norm = access_type(d.get("TIPO"))
        sign = -1 if tipo_norm == "NC" else 1
        numero = int(d.get("Numero") or 0)
        rows.append(
            {
                "origen": "ACCESS",
                "fecha": d.get("FECHA", "")[:10],
                "tipo": d.get("TIPO"),
                "tipo_norm": tipo_norm,
                "punto_venta": "",
                "numero": numero,
                "numero_texto": str(numero),
                "emisor": d.get("PROVEEDOR", ""),
                "cuit": "",
                "iva": sign * n(d.get("IVA")),
                "iva_original": n(d.get("IVA")),
                "importe": sign * n(d.get("IMPORTE")),
                "id_gasto": d.get("IdGASTO"),
                "user_id": d.get("UserID"),
            }
        )
    return rows


def key(row):
    return (row["tipo_norm"], row["numero"])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--arca", required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--month", type=int, required=True)
    parser.add_argument("--out", default="")
    args = parser.parse_args()

    arca = load_arca(Path(args.arca), args.year, args.month)
    access = access_to_rows(run_access_query(args.year, args.month))

    arca_by_key = defaultdict(list)
    access_by_key = defaultdict(list)
    for row in arca:
        arca_by_key[key(row)].append(row)
    for row in access:
        access_by_key[key(row)].append(row)

    all_keys = sorted(set(arca_by_key) | set(access_by_key))
    diffs = []
    matched = 0
    for k in all_keys:
        a_rows = arca_by_key.get(k, [])
        x_rows = access_by_key.get(k, [])
        count = max(len(a_rows), len(x_rows))
        for i in range(count):
            a = a_rows[i] if i < len(a_rows) else None
            x = x_rows[i] if i < len(x_rows) else None
            if a and x:
                matched += 1
                iva_diff = x["iva"] - a["iva"]
                total_diff = x["importe"] - a["importe"]
                status = "OK" if abs(iva_diff) <= 1 and abs(total_diff) <= 1 else "DIFERENCIA_IMPORTE"
            elif a:
                iva_diff = -a["iva"]
                total_diff = -a["importe"]
                status = "FALTA_EN_ACCESS"
            else:
                iva_diff = x["iva"]
                total_diff = x["importe"]
                status = "SOBRA_EN_ACCESS"
            if status != "OK":
                diffs.append(
                    {
                        "Estado": status,
                        "Tipo": (a or x)["tipo_norm"],
                        "Numero": (a or x)["numero"],
                        "ARCA_Fecha": a["fecha"] if a else "",
                        "ARCA_Emisor": a["emisor"] if a else "",
                        "ARCA_CUIT": a["cuit"] if a else "",
                        "ARCA_Moneda": a["moneda"] if a else "",
                        "ARCA_TipoCambio": round(a["tipo_cambio"], 4) if a else "",
                        "ARCA_IVA": round(a["iva"], 2) if a else "",
                        "ARCA_Total": round(a["importe"], 2) if a else "",
                        "Access_Fecha": x["fecha"] if x else "",
                        "Access_Proveedor": x["emisor"] if x else "",
                        "Access_IdGASTO": x.get("id_gasto", "") if x else "",
                        "Access_UserID": x.get("user_id", "") if x else "",
                        "Access_IVA": round(x["iva"], 2) if x else "",
                        "Access_Total": round(x["importe"], 2) if x else "",
                        "Diferencia_IVA_Access_menos_ARCA": round(iva_diff, 2),
                        "Diferencia_Total_Access_menos_ARCA": round(total_diff, 2),
                    }
                )

    out = Path(args.out) if args.out else ROOT / "outputs" / f"comparacion_arca_access_gastos_{args.year}-{args.month:02d}.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(diffs[0].keys()) if diffs else ["Estado"])
        writer.writeheader()
        writer.writerows(diffs)

    arca_iva = sum(r["iva"] for r in arca)
    access_iva = sum(r["iva"] for r in access)
    summary = {
        "periodo": f"{args.year}-{args.month:02d}",
        "arca_comprobantes": len(arca),
        "access_comprobantes": len(access),
        "coincidencias_por_tipo_numero": matched,
        "diferencias": len(diffs),
        "arca_iva": money(arca_iva),
        "access_iva": money(access_iva),
        "diferencia_iva_access_menos_arca": money(access_iva - arca_iva),
        "detalle_csv": str(out),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
